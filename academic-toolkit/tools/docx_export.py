# -*- coding: utf-8 -*-
'''Markdown → DOCX 转换引擎（中文学术文档格式）

用法：
    python docx_export.py --source paper.md --output paper.docx
    python docx_export.py --source paper.md --output paper.docx --profile literature_review.json
    python docx_export.py --source paper.md --output paper.docx --workspace /path/to/workspace

依赖：python-docx（已在 requirements.txt 中）
'''
from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Pt, RGBColor

log = logging.getLogger(__name__)

def _resolve_profiles_dir() -> Path:
    """样式 profile 目录解析（2026-09-23 v2.0：原 tools/docx_style_profiles/ 已迁
    third_party/docx-style-profiles/，故按候选列表探测，兼容两种布局）。"""
    base = Path(__file__).resolve().parent
    candidates = [
        base / 'docx_style_profiles',                               # 旧布局（本机历史）
        base.parent / 'third_party' / 'docx-style-profiles',        # v2.0 新布局
    ]
    for c in candidates:
        if c.is_dir():
            return c
    return candidates[0]


_PROFILES_DIR = _resolve_profiles_dir()
_DEFAULT_PROFILE = 'default_cn_thesis.json'


def load_style_profile(profile_path: Optional[Path] = None) -> dict[str, Any]:
    '''加载样式配置 JSON。缺失字段用默认值填充。'''
    default_path = _PROFILES_DIR / _DEFAULT_PROFILE
    if not default_path.exists():
        log.warning('Default profile not found: %s', default_path)
        return _builtin_defaults()
    default_data = json.loads(default_path.read_text(encoding='utf-8'))
    if profile_path is None or not profile_path.exists():
        return default_data
    try:
        custom_data = json.loads(profile_path.read_text(encoding='utf-8'))
    except (json.JSONDecodeError, OSError) as e:
        log.warning('Failed to load profile %s: %s, using defaults', profile_path, e)
        return default_data
    return _deep_merge(default_data, custom_data)


def _builtin_defaults() -> dict[str, Any]:
    '''内置兜底默认值（当 JSON 文件不存在时使用）。'''
    return {
        'page': {
            'size': 'A4',
            'margin_top_cm': 2.5,
            'margin_bottom_cm': 2.5,
            'margin_left_cm': 2.5,
            'margin_right_cm': 2.5,
        },
        'fonts': {
            'chinese_heading': 'SimHei',
            'chinese_body': 'SimSun',
            'latin': 'Times New Roman',
            'monospace': 'Consolas',
        },
        'headings': {
            'level1_pt': 18,
            'level2_pt': 15,
            'level3_pt': 12,
            'bold': True,
            'level1_alignment': 'center',
            'level2_alignment': 'left',
            'level3_alignment': 'left',
            'level1_page_break_before': True,
            'level2_page_break_before': False,
            'level3_page_break_before': False,
        },
        'body': {
            'font_size_pt': 12,
            'line_spacing': 1.5,
            'first_line_indent_chars': 2,
            'space_before_pt': 0,
            'space_after_pt': 0,
        },
        'title': {
            'font_size_pt': 18,
            'bold': True,
            'alignment': 'center',
            'font_family': 'SimHei',
        },
        'table': {
            'top_border_pt': 1.5,
            'header_border_pt': 0.75,
            'bottom_border_pt': 1.5,
            'font_size_pt': 10.5,
            'header_bold': True,
            'cell_alignment': 'center',
        },
        'references': {
            'hanging_indent_cm': 0.74,
            'font_size_pt': 10.5,
            'numbering_style': 'bracket',
        },
        'image': {
            'max_width_cm': 14,
            'max_height_cm': 20,
            'alignment': 'center',
        },
        'code_block': {
            'font_size_pt': 9,
            'line_spacing': 1,
            'background_color': 'F5F5F5',
        },
    }


def _deep_merge(base: dict, override: dict) -> dict:
    '''递归合并字典，override 覆盖 base。'''
    merged = dict(base)
    for key, value in override.items():
        if key in merged and isinstance(merged[key], dict) and isinstance(value, dict):
            merged[key] = _deep_merge(merged[key], value)
            continue
        merged[key] = value
    return merged


def _read_markdown(path: Path) -> str:
    '''读取 Markdown 文件，尝试多种编码。'''
    raw = path.read_bytes()
    for enc in ('utf-8', 'utf-8-sig', 'gbk', 'gb2312', 'gb18030', 'latin-1'):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError(f'无法解码文件 {path}：尝试了 UTF-8, GBK, Latin-1 均失败')


def _ensure_png_for_docx(md_content: str, workspace: Path) -> None:
    '''扫描 markdown 中的图片引用，若 .png 不存在但同名 .pdf 存在，自动转换。
    优先使用 PyMuPDF (fitz)，回退到 pdf2image。失败则保留原始情况，由后续渲染流程提示。

    ⛔ 主动兜底: 不仅扫 markdown 引用, 还把 figures/ 下所有 fig_*.pdf 强制转成同名
    png (没有同名 png 时才转), 覆盖以下情况:
    - AI 在 markdown 里写错路径 (例如写 figures/q1.pdf 而非 figures/fig_q1.pdf)
    - AI 把图引用写成 \\includegraphics 形式 (没 ![]()) 导致引用扫描漏掉
    - 写作步漏转 PDF 但产物本身在 figures/ 下完整存在
    '''
    refs = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', md_content)
    for ref in refs:
        ref_clean = ref.strip().split(' ')[0]
        if not ref_clean:
            continue
        suffix = Path(ref_clean).suffix.lower()
        if suffix not in ('.png', '.jpg', '.jpeg'):
            continue
        cand_paths = [workspace / ref_clean, workspace / 'figures' / Path(ref_clean).name]
        png_path = next((p for p in cand_paths if p.exists()), None)
        if png_path is not None:
            continue
        pdf_candidates = [
            workspace / Path(ref_clean).with_suffix('.pdf'),
            workspace / 'figures' / Path(ref_clean).with_suffix('.pdf').name,
        ]
        pdf_path = next((p for p in pdf_candidates if p.exists()), None)
        if pdf_path is None:
            continue
        target_png = pdf_path.with_suffix('.png')
        if target_png.exists():
            continue
        if _convert_pdf_to_png(pdf_path, target_png, dpi=350):
            log.info('Auto-converted %s -> %s for Word embedding (md ref)', pdf_path.name, target_png.name)
            _trim_png_whitespace(target_png)
        else:
            log.warning('Could not convert %s to PNG for md ref', pdf_path.name)

    figures_dir = workspace / 'figures'
    if not figures_dir.is_dir():
        return
    for pdf in figures_dir.glob('fig_*.pdf'):
        target_png = pdf.with_suffix('.png')
        if target_png.exists():
            continue
        if _convert_pdf_to_png(pdf, target_png, dpi=350):
            log.info('Proactive PDF→PNG for %s -> %s', pdf.name, target_png.name)
            _trim_png_whitespace(target_png)
        else:
            log.warning('Proactive PDF→PNG failed for %s (no PyMuPDF / pdf2image)', pdf.name)


def _convert_pdf_to_png(pdf_path: Path, png_path: Path, dpi: int = 350) -> bool:
    '''把单页 PDF 转成 PNG。优先 PyMuPDF(fitz)，回退 pdf2image。成功返回 True。'''
    try:
        import fitz
        doc = fitz.open(str(pdf_path))
        if doc.page_count > 0:
            pix = doc[0].get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
            pix.save(str(png_path))
            doc.close()
            return True
        doc.close()
    except Exception as e:
        log.debug('fitz convert failed for %s: %s', pdf_path.name, e)

    try:
        from pdf2image import convert_from_path
        images = convert_from_path(str(pdf_path), dpi=dpi, first_page=1, last_page=1)
        if images:
            images[0].save(str(png_path), 'PNG')
            return True
    except Exception as e:
        log.debug('pdf2image convert failed for %s: %s', pdf_path.name, e)

    return False


def _ensure_tikz_png_for_docx(workspace: Path) -> None:
    '''TikZ 兜底：把 figures/tikz_*.pdf 转成 PNG 供 Word 嵌入。

    TikZ 由 paper-figure-drawio 编译成 figures/tikz_diagrams.pdf（多图则多页）。
    Word 不能嵌 PDF，必须转 PNG。处理两种情况：
    - 单页 PDF → 同名 PNG（tikz_diagrams.pdf → tikz_diagrams.png）
    - 多页 PDF → 逐页 PNG（tikz_diagrams.pdf → tikz_diagrams_1.png / _2.png ...）

    幂等：目标 PNG 已存在则跳过。不删除原 PDF（保留供 PDF 模式/排查），
    Word 嵌入时 _add_image 会优先用 PNG。
    '''
    figures_dir = workspace / 'figures'
    if not figures_dir.is_dir():
        return
    tikz_pdfs = sorted(figures_dir.glob('tikz_*.pdf'))
    if not tikz_pdfs:
        return

    try:
        import fitz
        _has_fitz = True
    except Exception:
        _has_fitz = False

    for tpdf in tikz_pdfs:
        bn = tpdf.stem
        page_count = 1
        if _has_fitz:
            try:
                import fitz
                _d = fitz.open(str(tpdf))
                page_count = _d.page_count
                _d.close()
            except Exception:
                page_count = 1
        if page_count <= 1:
            target = figures_dir / f'{bn}.png'
            if not target.exists() and _convert_pdf_to_png(tpdf, target):
                log.info('TikZ: converted %s -> %s for Word', tpdf.name, target.name)
                _trim_png_whitespace(target)
            continue
        try:
            import fitz
            _d = fitz.open(str(tpdf))
            for i in range(_d.page_count):
                target = figures_dir / f'{bn}_{i + 1}.png'
                if target.exists():
                    continue
                pix = _d[i].get_pixmap(matrix=fitz.Matrix(4.861111111111111, 4.861111111111111), alpha=False)
                pix.save(str(target))
                log.info('TikZ: converted %s page %d -> %s', tpdf.name, i + 1, target.name)
                _trim_png_whitespace(target)
            _d.close()
            first_png = figures_dir / f'{bn}.png'
            if not first_png.exists():
                _convert_pdf_to_png(tpdf, first_png)
                if first_png.exists():
                    _trim_png_whitespace(first_png)
        except Exception as e:
            log.warning('TikZ multi-page conversion failed for %s: %s', tpdf.name, e)
            continue
    return


def _trim_png_whitespace(png_path: Path, bg_threshold: int = 240, padding: int = 12) -> bool:
    '''裁掉 PNG 四周的白边。

    Args:
        png_path: 目标 PNG 路径
        bg_threshold: 背景判定阈值（亮度 ≥ bg_threshold 视为白边，容忍轻微抗锯齿灰）
        padding: 裁完后保留的像素边距（避免贴边）

    Returns:
        True 表示成功裁边并保存；False 表示跳过（PIL 不可用或图无白边）
    '''
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return False
    try:
        im = Image.open(str(png_path))
        if im.mode in ('RGBA', 'LA'):
            bg = Image.new('RGB', im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        elif im.mode != 'RGB':
            im = im.convert('RGB')
        bg_ref = Image.new('RGB', im.size, (255, 255, 255))
        diff = ImageChops.difference(im, bg_ref)
        bbox_data = diff.point(lambda x: 0 if x < (255 - bg_threshold) else 255)
        bbox = bbox_data.getbbox()
        if bbox is None:
            return False
        x0, y0, x1, y1 = bbox
        orig_w, orig_h = im.size

        x0 = max(0, x0 - padding)
        y0 = max(0, y0 - padding)
        x1 = min(orig_w, x1 + padding)
        y1 = min(orig_h, y1 + padding)
        new_w, new_h = x1 - x0, y1 - y0

        cut_w_pct = (orig_w - new_w) / orig_w
        cut_h_pct = (orig_h - new_h) / orig_h
        content_area_pct = (new_w * new_h) / (orig_w * orig_h)
        should_trim = cut_w_pct > 0.02 or cut_h_pct > 0.02 or content_area_pct < 0.5
        if should_trim:
            cropped = im.crop((x0, y0, x1, y1))
            cropped.save(str(png_path), 'PNG', optimize=True)
            log.info(
                'Trimmed %s whitespace: %dx%d -> %dx%d (content %.0f%% -> 100%%)',
                png_path.name, orig_w, orig_h, new_w, new_h, content_area_pct * 100,
            )
            return True
        return False
    except Exception as e:
        log.debug('Trim whitespace failed for %s: %s', png_path.name, e)
        return False


def _set_run_fonts(run, profile: dict[str, Any], style_type: str) -> None:
    '''为 run 设置中英文字体。style_type: body / heading / code / reference'''
    fonts = profile.get('fonts', {})
    if style_type == 'heading':
        cn_font = fonts.get('chinese_heading', 'SimHei')
        en_font = fonts.get('latin', 'Times New Roman')
    elif style_type == 'code':
        cn_font = fonts.get('monospace', 'Consolas')
        en_font = fonts.get('monospace', 'Consolas')
    elif style_type == 'reference':
        cn_font = fonts.get('chinese_body', 'SimSun')
        en_font = fonts.get('latin', 'Times New Roman')
    else:
        cn_font = fonts.get('chinese_body', 'SimSun')
        en_font = fonts.get('latin', 'Times New Roman')
    run.font.name = en_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn('w:rFonts'))
    if r_fonts is None:
        from docx.oxml import OxmlElement
        r_fonts = OxmlElement('w:rFonts')
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn('w:eastAsia'), cn_font)
    r_fonts.set(qn('w:ascii'), en_font)
    r_fonts.set(qn('w:hAnsi'), en_font)


_ALIGNMENT_MAP = {
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _clean_markdown_inline(text: str) -> str:
    '''清理 Markdown 行内标记（加粗/斜体/代码/链接），保留纯文本。'''
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    # 单星号斜体（在加粗处理之后）
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'(?<!\w)_(.+?)_(?!\w)', r'\1', text)
    # 行内代码
    text = re.sub(r'`(.+?)`', r'\1', text)
    # 链接 [text](url)
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
    return text.strip()


def _add_heading(doc: Document, text: str, level: int, profile: dict, first_heading: bool = False):
    """添加标题段落。

    first_heading=True（第一个 H1）走 profile['title'] 样式（论文标题，居中、加粗、22pt）；
    其他标题走 profile['headings'] 的 levelN 配置。
    """
    headings_cfg = profile.get('headings', {})
    title_cfg = profile.get('title', {})
    if first_heading and level == 1 and title_cfg:
        font_size = title_cfg.get('font_size_pt', 22)
        bold = title_cfg.get('bold', True)
        alignment = title_cfg.get('alignment', 'center')
        page_break = False
    else:
        size_key = f'level{level}_pt'
        if size_key in headings_cfg:
            font_size = headings_cfg[size_key]
        else:
            font_size = 12
            for _try_lvl in range(level - 1, 0, -1):
                _try_key = f'level{_try_lvl}_pt'
                if _try_key in headings_cfg:
                    font_size = max(headings_cfg[_try_key] - level - _try_lvl, 10.5)
                    break
        bold = headings_cfg.get('bold', True)
        align_key = f'level{level}_alignment'
        alignment = headings_cfg.get(align_key, 'left')
        page_break_key = f'level{level}_page_break_before'
        page_break = headings_cfg.get(page_break_key, False)
    para = doc.add_paragraph()
    if page_break and not first_heading:
        from docx.oxml import OxmlElement
        pPr = para._p.get_or_add_pPr()
        pb = OxmlElement('w:pageBreakBefore')
        pPr.append(pb)
    para.alignment = _ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    if first_heading and level == 1 and title_cfg:
        para.paragraph_format.space_before = Pt(24)
        para.paragraph_format.space_after = Pt(18)
    run = para.add_run(_clean_markdown_inline(text))
    run.bold = bold
    run.font.size = Pt(font_size)
    _set_run_fonts(run, profile, 'heading')
    return para


def _add_body_paragraph(doc: Document, text: str, profile: dict):
    '''添加正文段落（首行缩进 + 中文字体）。'''
    body_cfg = profile.get('body', {})
    font_size = body_cfg.get('font_size_pt', 12)
    line_spacing = body_cfg.get('line_spacing', 1.5)
    indent_chars = body_cfg.get('first_line_indent_chars', 2)
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = line_spacing
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    # 首行缩进 2 字符 ≈ 2 × 字号（pt）
    indent_pt = indent_chars * font_size
    para.paragraph_format.first_line_indent = Pt(indent_pt)
    para.paragraph_format.space_before = Pt(body_cfg.get('space_before_pt', 0))
    para.paragraph_format.space_after = Pt(body_cfg.get('space_after_pt', 0))
    run = para.add_run(_clean_markdown_inline(text))
    run.font.size = Pt(font_size)
    _set_run_fonts(run, profile, 'body')
    return para


def _add_runs_with_inline_bold(para, text: str, profile: dict, role: str, font_size: float):
    '''把 text 按 **加粗** 拆成多个 run：加粗段 bold=True，其余段正常；
    每段都先 _clean_markdown_inline 清掉其它行内标记（斜体/代码/链接），再逐 run 套字体+字号。
    仅用于摘要正文（竞赛摘要加粗关键方法/结果）。普通正文不走此路径，行为不变。

    边缘：
    - 无 **加粗** → parts 仅一段非加粗，等价于旧的单 run（清洗后加一个 run）。
    - 连续/多处加粗、加粗跨标点 → re.split 天然正确切分。
    - 未闭合 **（如 "**未闭合"）→ 不成对，不匹配加粗，整段当普通文本清洗（与旧行为一致）。
    - 加粗段清洗后为空 → 跳过该 run，不产生空粗体。
    - 全部清空（理论不会）→ 兜底加一个空 run 保住段落存在。
    '''
    parts = re.split(r'(\*\*.+?\*\*)', text)
    any_run = False
    for seg in parts:
        if not seg:
            continue
        m = re.match(r'^\*\*(.+?)\*\*$', seg)
        if m:
            content = _clean_markdown_inline(m.group(1))
            is_bold = True
        else:
            content = _clean_markdown_inline(seg)
            is_bold = False
        if not content:
            continue
        run = para.add_run(content)
        run.bold = is_bold
        run.font.size = Pt(font_size)
        _set_run_fonts(run, profile, role)
        any_run = True
    if not any_run:
        run = para.add_run('')
        run.font.size = Pt(font_size)
        _set_run_fonts(run, profile, role)
    return para


def _add_abstract_body_paragraph(doc: Document, text: str, profile: dict):
    '''摘要正文段落：段落格式与正文完全一致（首行缩进/行距/字体），
    但 run 层支持 **加粗** 行内解析，用于竞赛摘要加粗关键方法名/关键结果数值。'''
    body_cfg = profile.get('body', {})
    font_size = body_cfg.get('font_size_pt', 12)
    line_spacing = body_cfg.get('line_spacing', 1.5)
    indent_chars = body_cfg.get('first_line_indent_chars', 2)
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = line_spacing
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    indent_pt = indent_chars * font_size
    para.paragraph_format.first_line_indent = Pt(indent_pt)
    para.paragraph_format.space_before = Pt(body_cfg.get('space_before_pt', 0))
    para.paragraph_format.space_after = Pt(body_cfg.get('space_after_pt', 0))
    _add_runs_with_inline_bold(para, text, profile, 'body', font_size)
    return para


def _add_reference_paragraph(doc: Document, text: str, profile: dict):
    '''添加参考文献段落（悬挂缩进）。'''
    ref_cfg = profile.get('references', {})
    font_size = ref_cfg.get('font_size_pt', 10.5)
    hanging_cm = ref_cfg.get('hanging_indent_cm', 0.74)
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(-hanging_cm)
    para.paragraph_format.left_indent = Cm(hanging_cm)
    para.paragraph_format.line_spacing = 1.25
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(_clean_markdown_inline(text))
    run.font.size = Pt(font_size)
    _set_run_fonts(run, profile, 'reference')
    return para


def _add_abstract_heading(doc: Document, text: str, profile: dict):
    '''添加摘要标题（居中、加粗，对标 cumcmthesis.cls 的 \\zihao{4}\\bfseries 摘要）。'''
    abstract_cfg = profile.get('abstract', {})
    label_size = abstract_cfg.get('label_size_pt', 14)
    label_bold = abstract_cfg.get('label_bold', True)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(_clean_markdown_inline(text))
    run.bold = label_bold
    run.font.size = Pt(label_size)
    _set_run_fonts(run, profile, 'heading')
    return para


def _add_keywords_paragraph(doc: Document, label: str, value: str, profile: dict):
    '''添加关键词段落（黑体加粗标签 + 正文内容）。'''
    kw_cfg = profile.get('keywords', {})
    body_cfg = profile.get('body', {})
    label_size = kw_cfg.get('label_size_pt', body_cfg.get('font_size_pt', 12))
    label_bold = kw_cfg.get('label_bold', True)
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(label_size * 2)
    para.paragraph_format.line_spacing = body_cfg.get('line_spacing', 1.5)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run_label = para.add_run(label)
    run_label.bold = label_bold
    run_label.font.size = Pt(label_size)
    _set_run_fonts(run_label, profile, 'heading')
    run_value = para.add_run(_clean_markdown_inline(value))
    run_value.font.size = Pt(label_size)
    _set_run_fonts(run_value, profile, 'body')
    return para


def _parse_frontmatter(content: str) -> tuple[dict, str]:
    '''解析并丢弃 YAML frontmatter（封面元数据已不再渲染，但要确保它不会被当正文）。

    返回 (meta_dict, body_without_frontmatter)。
    '''
    m = re.match(r'^---\s*\r?\n([\s\S]*?)\r?\n---\s*\r?\n', content)
    if not m:
        return {}, content
    meta = {}
    for line in m.group(1).split('\n'):
        kv = re.match(r'^([\w\-]+)\s*:\s*(.*)$', line)
        if kv:
            v = kv.group(2).strip()
            if (v.startswith('"') and v.endswith('"')) or (v.startswith("'") and v.endswith("'")):
                v = v[1:-1]
            meta[kv.group(1)] = v
    return meta, content[m.end():]


def _add_code_block(doc: Document, lines: list[str], profile: dict):
    '''添加代码块（等宽字体 + 浅灰底色 + 每行独立段落保留缩进与换行）。'''
    code_cfg = profile.get('code_block', {})
    font_size = code_cfg.get('font_size_pt', 9)
    bg_color = code_cfg.get('background_color', 'F5F5F5')
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for line in lines:
        expanded = line.replace('\t', '    ').rstrip()
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.left_indent = Cm(0.5)
        para.paragraph_format.line_spacing = code_cfg.get('line_spacing', 1.0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        try:
            pPr = para._p.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), bg_color)
            pPr.append(shd)
        except Exception:
            pass
        run = para.add_run(expanded if expanded else ' ')
        run.font.size = Pt(font_size)
        _set_run_fonts(run, profile, 'code')
    return


def _add_image(doc: Document, image_ref: str, workspace: Optional[Path], profile: dict):
    '''嵌入图片或插入占位符。image_ref 是 Markdown 中的路径。

    宽度策略：正常图统一按 max_width_cm（≈文本宽 85%）嵌入，不大不小；
    真·低像素图（放到目标宽会糊，有效 DPI<150）才回退原生尺寸；
    瘦长图按 max_height_cm 封顶反向缩宽，防止溢出页面。
    '''
    img_cfg = profile.get('image', {})
    target_width_cm = float(img_cfg.get('max_width_cm', 14))
    max_height_cm = float(img_cfg.get('max_height_cm', 20))
    col_w = _text_column_width_cm(profile)
    if col_w < target_width_cm:
        target_width_cm = max(col_w - 0.2, 2.0)
    alignment = img_cfg.get('alignment', 'center')

    image_path = None
    if workspace:
        candidate = workspace / image_ref
        if candidate.exists():
            image_path = candidate
        else:
            candidate2 = workspace / 'figures' / Path(image_ref).name
            if candidate2.exists():
                image_path = candidate2

    if image_path and image_path.suffix.lower() in ('.pdf', '.svg'):
        png_alt = image_path.with_suffix('.png')
        if png_alt.exists():
            image_path = png_alt

    if image_path is None and workspace:
        ref_path = Path(image_ref)
        if ref_path.suffix.lower() in ('.png', '.jpg', '.jpeg'):
            for base in (workspace / ref_path.with_suffix('.pdf'),
                         workspace / 'figures' / ref_path.with_suffix('.pdf').name):
                if base.exists():
                    log.warning('Image %s not found, but %s exists. Word 需要 PNG。', image_ref, base)
                    break

    if image_path and image_path.exists():
        ext = image_path.suffix.lower()
        if ext in ('.svg', '.pdf'):
            log.warning('%s not supported by python-docx, skipping: %s', ext, image_ref)
            para = doc.add_paragraph()
            para.alignment = _ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
            run = para.add_run(f'[{ext.upper().lstrip(".")}图片无法嵌入 Word，需要 PNG: {image_ref}]')
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(153, 153, 153)
            return None
        para = doc.add_paragraph()
        para.alignment = _ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
        para.paragraph_format.first_line_indent = Pt(0)
        run = para.add_run()
        try:
            actual_width_cm = target_width_cm
            try:
                from PIL import Image as _PILImage
                with _PILImage.open(image_path) as _im:
                    img_w_px, img_h_px = _im.size
                eff_dpi = img_w_px / (target_width_cm / 2.54)
                if eff_dpi < 150:
                    native_cm = img_w_px / 96 * 2.54
                    actual_width_cm = min(native_cm, target_width_cm)
                if img_w_px > 0:
                    aspect = img_h_px / img_w_px
                    height_cm = actual_width_cm * aspect
                    if height_cm > max_height_cm:
                        actual_width_cm = max_height_cm / aspect
            except Exception:
                pass
            run.add_picture(str(image_path), width=Cm(actual_width_cm))
            return None
        except Exception as e:
            log.warning('Failed to embed image %s: %s', image_path, e)
            para.clear()
            run2 = para.add_run(f'[图片加载失败: {image_ref}]')
            run2.font.size = Pt(10)
            run2.font.color.rgb = RGBColor(153, 153, 153)
            return None

    para = doc.add_paragraph()
    para.alignment = _ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
    para.paragraph_format.first_line_indent = Pt(0)
    run = para.add_run(f'[图片缺失: {image_ref}]')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(153, 153, 153)
    return None


def _add_three_line_table(doc: Document, headers: list[str], rows: list[list[str]], profile: dict):
    '''添加三线表。'''
    table_cfg = profile.get('table', {})
    font_size = table_cfg.get('font_size_pt', 10.5)
    header_bold = table_cfg.get('header_bold', True)
    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        run = para.add_run(_clean_markdown_inline(h))
        run.bold = header_bold
        run.font.size = Pt(font_size)
        _set_run_fonts(run, profile, 'body')

    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= col_count:
                continue
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            run = para.add_run(_clean_markdown_inline(cell_text))
            run.font.size = Pt(font_size)
            _set_run_fonts(run, profile, 'body')

    _fix_table_width(table, profile, col_count)
    _apply_three_line_borders(table, profile)
    return table


def _fix_table_width(table, profile: dict, col_count: int):
    '''把表宽锁到可用文本宽、列宽平分、固定布局。双栏时 = 单栏栏宽，防溢出。'''
    if col_count <= 0:
        return None
    avail_twips = int(_text_column_width_cm(profile) * 567)
    per_col = max(avail_twips // col_count, 1)
    total_twips = per_col * col_count
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ('w:tblW', 'w:tblLayout'):
        for existing in tblPr.findall(qn(tag)):
            tblPr.remove(existing)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_twips}" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc in grid.findall(qn('w:gridCol')):
            grid.remove(gc)
        for _ in range(col_count):
            grid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{per_col}"/>'))
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for existing in tcPr.findall(qn('w:tcW')):
                tcPr.remove(existing)
            tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{per_col}" w:type="dxa"/>'))
    return None


def _apply_three_line_borders(table, profile: dict):
    '''为表格应用三线表边框样式。'''
    table_cfg = profile.get('table', {})
    top_sz = int(table_cfg.get('top_border_pt', 1.5) * 8)
    header_sz = int(table_cfg.get('header_border_pt', 0.75) * 8)
    bottom_sz = int(table_cfg.get('bottom_border_pt', 1.5) * 8)
    row_count = len(table.rows)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            existing = tcPr.find(qn('w:tcBorders'))
            if existing is not None:
                tcPr.remove(existing)
            if row_idx == 0:
                borders_xml = (
                    f'<w:tcBorders {nsdecls("w")}'
                    f'>  <w:top w:val="single" w:sz="{top_sz}" w:space="0" w:color="000000"/>'
                    f'  <w:bottom w:val="single" w:sz="{header_sz}" w:space="0" w:color="000000"/>'
                    f'  <w:left w:val="nil"/>  <w:right w:val="nil"/></w:tcBorders>'
                )
            elif row_idx == row_count - 1:
                borders_xml = (
                    f'<w:tcBorders {nsdecls("w")}'
                    f'>  <w:top w:val="nil"/>'
                    f'  <w:bottom w:val="single" w:sz="{bottom_sz}" w:space="0" w:color="000000"/>'
                    f'  <w:left w:val="nil"/>  <w:right w:val="nil"/></w:tcBorders>'
                )
            else:
                borders_xml = (
                    f'<w:tcBorders {nsdecls("w")}'
                    f'>  <w:top w:val="nil"/>  <w:bottom w:val="nil"/>'
                    f'  <w:left w:val="nil"/>  <w:right w:val="nil"/></w:tcBorders>'
                )
            tcPr.append(parse_xml(borders_xml))
    return None


_CN_SECTION_NUM_RE = re.compile('^(?:第[一二三四五六七八九十百千零〇\\d]+[章节部篇]|[一二三四五六七八九十]+、|[（(][一二三四五六七八九十]+[)）])')

_NUMERIC_SECTION_RE = re.compile('^\\d+(?:\\.\\d+){0,4}\\.?\\s')

_KNOWN_HEADING_LABELS = ('摘要', 'Abstract', '关键词', 'Keywords', 'Key words', '参考文献', 'References', '致谢', 'Acknowledgement', 'Acknowledgements', '附录', 'Appendix')

_FIG_TABLE_CAPTION_RE = re.compile('^(?:(?:图|表)\\s*[\\dA-Za-z][\\dA-Za-z\\-\\.]*\\s*[：:、\\.\\-]|(?:图|表)\\s*[：:]|(?:Figure|Fig\\.?|Table|Tab\\.?)\\s*[\\dA-Za-z][\\dA-Za-z\\-\\.]*)')

_BARE_EQ_NUMBER_RE = re.compile('^\\s*[\\[\\(【]\\s*\\d+(?:\\.\\d+)?\\s*[\\]\\)】]\\s*$')


def _classify_pseudo_heading(text: str) -> int:
    '''如果 text 看起来像被误写成 list 的章节标题，返回应有的级别（1/2/3/4）；否则 0。'''
    s = text.strip()
    if not s or len(s) > 80:
        return 0
    for lbl in _KNOWN_HEADING_LABELS:
        if (s == lbl or s.startswith(lbl + '（') or s.startswith(lbl + '(')
                or s.startswith(lbl + ':') or s.startswith(lbl + '：')):
            return 2
    if _CN_SECTION_NUM_RE.match(s):
        return 1
    m = _NUMERIC_SECTION_RE.match(s)
    if m:
        prefix = m.group(0).rstrip('. ').rstrip()
        depth = prefix.count('.') + 1
        return min(depth, 4)
    return 0


def _normalize_pseudo_markup(content: str) -> str:
    '''把 LLM 误写为 list 的章节标题、图注、独行公式编号还原成合法 markdown。

    不依赖语义判断，只看结构特征：
      - "- 一、问题重述"          → "# 一、问题重述"
      - "- 1.1 问题背景"          → "## 1.1 问题背景"
      - "- 1.1.1 子小节"          → "### 1.1.1 子小节"
      - "- 摘要"                  → "## 摘要"
      - "- 图 4-1：方法对比结果"  → "图 4-1：方法对比结果"  (去 bullet, 保留正文)
      - "- (1)" / "- [1]"         → 与上一行公式合并成 "$$...$$ (1)"
    在 ``` 代码块内不处理，避免破坏代码示例。
    '''
    if not content:
        return content
    out = []
    in_code = False
    fence_re = re.compile('^\\s*```')
    list_re = re.compile('^(\\s*)([-*+])\\s+(.+?)\\s*$')
    lines = content.splitlines()
    for raw in lines:
        if fence_re.match(raw):
            in_code = not in_code
            out.append(raw)
            continue
        if in_code:
            out.append(raw)
            continue
        m = list_re.match(raw)
        if m:
            indent = m.group(1)
            body = m.group(3)
            if indent == '':
                lvl = _classify_pseudo_heading(body)
                if lvl > 0:
                    out.append('#' * lvl + ' ' + body)
                    continue
                if _FIG_TABLE_CAPTION_RE.match(body):
                    out.append(body)
                    continue
                if _BARE_EQ_NUMBER_RE.match(body):
                    j = len(out) - 1
                    while j >= 0 and not out[j].strip():
                        j -= 1
                    if j >= 0:
                        prev = out[j].rstrip()
                        if prev.endswith('$$') or prev.count('$') >= 2:
                            number = body.strip()
                            num_match = re.search(r'\d+(?:\.\d+)?', number)
                            if num_match:
                                num = num_match.group(0)
                                out[j] = prev + f' ({num})'
                                continue
                    out.append(body)
                    continue
        out.append(raw)
    return '\n'.join(out) + ('\n' if content.endswith('\n') else '')


def markdown_to_docx(
    source_md: Path,
    output_path: Path,
    style_profile: Optional[Path] = None,
    workspace: Optional[Path] = None,
    engine: str = 'auto',
) -> Path:
    '''
    将 Markdown 文件转换为格式化的 .docx 文档。

    Args:
        source_md: 源 .md 文件路径
        output_path: 输出 .docx 文件路径
        style_profile: 可选的样式配置 JSON 路径
        workspace: 工作区根目录（用于解析图片相对路径）
        engine: 引擎选择 - "auto" / "python" / "node"
                auto 模式下：含 $...$ 公式时优先用 node，否则用 python

    Returns:
        生成的 .docx 文件路径

    Raises:
        FileNotFoundError: source_md 不存在
        ValueError: 文件无法解码
    '''
    if not source_md.exists():
        raise FileNotFoundError(f'源文件不存在: {source_md}')

    content = _read_markdown(source_md)

    # 预处理：把 LLM 误写成 list 的章节标题/图注/公式编号还原成合法 md
    _normalized_content = _normalize_pseudo_markup(content)
    if _normalized_content != content:
        content = _normalized_content
        # 写一个同目录临时文件供引擎使用
        _normalized_md = source_md.with_name(source_md.stem + '._normalized.md')
        _normalized_md.write_text(content, encoding='utf-8', newline='\n')
        _source_for_engine = _normalized_md
        log.info('Pseudo-markup normalized for %s', source_md.name)
    else:
        _source_for_engine = source_md

    if workspace is None:
        workspace = source_md.parent

    _ensure_png_for_docx(content, workspace)

    # TikZ 兜底：figures/tikz_*.pdf → PNG
    _ensure_tikz_png_for_docx(workspace)

    # 预裁剪 figures/ 下所有 png 白边
    figures_dir = workspace / 'figures'
    if figures_dir.is_dir():
        for png in figures_dir.glob('fig_*.png'):
            try:
                _trim_png_whitespace(png)
            except Exception as e:
                log.debug('Trim skipped for %s: %s', png.name, e)
        for png in figures_dir.glob('tikz_*.png'):
            try:
                _trim_png_whitespace(png)
            except Exception as e:
                log.debug('Trim skipped for %s: %s', png.name, e)

    chosen_engine = engine
    if chosen_engine == 'auto':
        # 检测复杂内容特征
        has_inline_math = bool(re.search('(?<![\\\\$])\\$[^$\\n]+\\$', content))
        has_block_math = bool(re.search('^\\s*\\$\\$', content, re.MULTILINE))
        has_table = bool(re.search('^\\s*\\|.+\\|\\s*$', content, re.MULTILINE)) and bool(
            re.search('^\\s*\\|[\\s\\-:|]+\\|\\s*$', content, re.MULTILINE))
        has_citation = bool(re.search('\\[\\d+(?:[,\\-]\\d+)*\\]', content))
        has_references = bool(re.search('^#{1,3}\\s*(?:参考文献|References)\\s*$', content, re.MULTILINE | re.IGNORECASE))
        has_academic = (bool(re.search('^#{1,3}\\s*(?:摘要|Abstract)\\s*$', content, re.MULTILINE | re.IGNORECASE))
                        or '关键词：' in content or 'Keywords:' in content)
        node_indicators = {
            '公式': has_inline_math or has_block_math,
            '三线表': has_table,
            '上标引用': has_citation,
            '参考文献': has_references,
            '学术格式': has_academic,
        }
        active_indicators = [k for k, v in node_indicators.items() if v]
        if active_indicators and _is_node_engine_available():
            chosen_engine = 'node'
            log.info('Auto-detected complex content [%s], using Node engine', ', '.join(active_indicators))
        elif active_indicators:
            log.warning(
                'Detected complex content [%s] but Node engine unavailable, falling back to Python (公式/三线表/上标可能不准确)',
                ', '.join(active_indicators),
            )
            chosen_engine = 'python'
        else:
            chosen_engine = 'python'
            log.info('Simple content, using Python engine')

    if chosen_engine == 'node':
        try:
            _result = _convert_with_node(_source_for_engine, output_path, workspace, style_profile)
        except Exception as e:
            log.warning('Node engine failed, falling back to Python: %s', e)
            chosen_engine = 'python'
        else:
            if _source_for_engine != source_md:
                try:
                    _source_for_engine.unlink()
                except Exception:
                    pass
            return _result

    try:
        return _convert_with_python(content, source_md, output_path, style_profile, workspace)
    finally:
        if _source_for_engine != source_md:
            try:
                _source_for_engine.unlink()
            except Exception:
                pass


def _is_node_engine_available() -> bool:
    '''检测 Node.js 引擎是否可用。'''
    engine_dir = _resolve_engine_dir()
    if not engine_dir or not engine_dir.is_dir():
        return False
    if not (engine_dir / 'md_to_docx.js').exists():
        return False
    if not (engine_dir / 'node_modules').is_dir():
        return False
    import shutil as _shutil
    if _shutil.which('node'):
        return True
    for candidate in ('D:\\nodejs\\node.exe', 'C:\\Program Files\\nodejs\\node.exe'):
        if Path(candidate).exists():
            return True
    return False


def _resolve_engine_dir() -> Optional[Path]:
    '''查找 docx-cn-engine 目录。优先级：
        1. 环境变量 DOCX_CN_ENGINE_DIR
        2. 与 docx_export.py 同目录下的 docx-cn-engine
        3. 上层 tools/docx-cn-engine（_utils/ 复制时实际目录）
    '''
    import os as _os
    env_path = _os.environ.get('DOCX_CN_ENGINE_DIR')
    if env_path:
        p = Path(env_path)
        if p.is_dir():
            return p
    same_dir = Path(__file__).resolve().parent / 'docx-cn-engine'
    if same_dir.is_dir():
        return same_dir
    cur = Path(__file__).resolve().parent
    for _ in range(5):
        cand = cur / 'tools' / 'docx-cn-engine'
        if cand.is_dir():
            return cand
        cur = cur.parent
    return None


def _find_node_executable() -> Optional[str]:
    '''查找 node 可执行文件路径。'''
    import shutil as _shutil
    found = _shutil.which('node')
    if found:
        return found
    for candidate in ('D:\\nodejs\\node.exe', 'C:\\Program Files\\nodejs\\node.exe'):
        if Path(candidate).exists():
            return candidate
    return None


def _convert_with_node(source_md: Path, output_path: Path, workspace: Optional[Path], style_profile: Optional[Path] = None) -> Path:
    '''使用 Node.js docx-cn-engine 转换。'''
    import subprocess as _subprocess
    node_exe = _find_node_executable()
    if not node_exe:
        raise RuntimeError('Node.js 不可用')
    engine_dir = _resolve_engine_dir()
    if engine_dir is None:
        raise RuntimeError('docx-cn-engine 目录找不到')
    script = engine_dir / 'md_to_docx.js'
    if not script.exists():
        raise RuntimeError(f'md_to_docx.js 不存在: {script}')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        node_exe, str(script),
        '--source', str(source_md),
        '--output', str(output_path),
    ]
    if workspace:
        cmd.extend(['--workspace', str(workspace)])
    if style_profile and style_profile.exists():
        cmd.extend(['--profile', str(style_profile)])
    result = _subprocess.run(
        cmd, capture_output=True, text=True, timeout=120, encoding='utf-8', errors='replace',
    )
    if result.returncode != 0:
        raise RuntimeError(f'Node 引擎失败 (exit {result.returncode}): {result.stderr or result.stdout}')
    log.info('DOCX exported via Node engine: %s (profile=%s)', output_path, style_profile.name if style_profile else 'default')
    return output_path


def _convert_with_python(content: str, source_md: Path, output_path: Path, style_profile: Optional[Path], workspace: Optional[Path]) -> Path:
    '''使用 python-docx 转换。'''
    profile = load_style_profile(style_profile)
    if workspace is None:
        workspace = source_md.parent
    # 丢弃 YAML frontmatter（封面元数据不渲染）
    fm_meta, body = _parse_frontmatter(content)
    lines = body.splitlines()

    doc = Document()
    _setup_page(doc, profile)
    _render_markdown(doc, lines, profile, workspace)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    log.info('DOCX exported via Python engine: %s (%d bytes)', output_path, output_path.stat().st_size)
    return output_path


def _setup_page(doc: Document, profile: dict):
    '''设置页面大小和页边距。'''
    page_cfg = profile.get('page', {})
    section = doc.sections[0]
    section.top_margin = Cm(page_cfg.get('margin_top_cm', 2.5))
    section.bottom_margin = Cm(page_cfg.get('margin_bottom_cm', 2.5))
    section.left_margin = Cm(page_cfg.get('margin_left_cm', 2.5))
    section.right_margin = Cm(page_cfg.get('margin_right_cm', 2.5))
    if page_cfg.get('size', 'A4') == 'A4':
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
    cols = int(page_cfg.get('columns', 1) or 1)
    if cols >= 2:
        sectPr = section._sectPr
        for existing in sectPr.findall(qn('w:cols')):
            sectPr.remove(existing)
        space_twips = int(float(page_cfg.get('column_spacing_cm', 0.6)) * 567)
        sectPr.append(parse_xml(f'<w:cols {nsdecls("w")} w:num="{cols}" w:space="{space_twips}" w:equalWidth="1"/>'))
        return None
    return None


def _text_column_width_cm(profile: dict) -> float:
    '''按边距+栏数算单栏可用文本宽（cm），供双栏时图/表缩放用。'''
    page_cfg = profile.get('page', {})
    left = float(page_cfg.get('margin_left_cm', 2.5))
    right = float(page_cfg.get('margin_right_cm', 2.5))
    cols = int(page_cfg.get('columns', 1) or 1)
    spacing = float(page_cfg.get('column_spacing_cm', 0.6))
    text_w = 21.0 - left - right
    if cols >= 2:
        return (text_w - spacing * (cols - 1)) / cols
    return text_w


def _render_markdown(doc: Document, lines: list[str], profile: dict, workspace: Path):
    '''逐行解析 Markdown 并渲染到 Document。'''
    i = 0
    total = len(lines)
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    in_references = False
    in_abstract = False
    first_h1_seen = False
    body_buffer = []

    def flush_body():
        if not body_buffer:
            return
        text = ' '.join(body_buffer)
        # 关键词行：加粗标签 + 内容
        kw_match = re.match('^\\*{0,2}(关键词|关键字|Key\\s*words|Keywords)\\*{0,2}\\s*[：:]\\s*(.+)$', text, re.IGNORECASE)
        if kw_match:
            label_raw = kw_match.group(1)
            value = kw_match.group(2)
            is_cn = any('一' <= c <= '鿿' for c in label_raw)
            label_out = '关键词：' if is_cn else 'Keywords: '
            _add_keywords_paragraph(doc, label_out, value, profile)
            body_buffer.clear()
            return
        if in_references and re.match('^\\[\\d+\\]', text):
            _add_reference_paragraph(doc, text, profile)
        elif in_abstract:
            # 摘要正文走 **加粗** 行内解析
            _add_abstract_body_paragraph(doc, text, profile)
        else:
            _add_body_paragraph(doc, text, profile)
        body_buffer.clear()

    def flush_table():
        if not table_lines:
            in_table = False
            return
        parsed_rows = []
        for tl in table_lines:
            cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if all(re.fullmatch('[-:]+', c) for c in cells):
                continue
            parsed_rows.append(cells)
        if len(parsed_rows) >= 2:
            headers = parsed_rows[0]
            data_rows = parsed_rows[1:]
            max_cols = max(len(headers), max((len(r) for r in data_rows), default=0))
            headers = headers + [''] * (max_cols - len(headers))
            data_rows = [r + [''] * (max_cols - len(r)) for r in data_rows]
            _add_three_line_table(doc, headers, data_rows, profile)
        elif len(parsed_rows) == 1:
            _add_body_paragraph(doc, ' | '.join(parsed_rows[0]), profile)
        table_lines.clear()
        return

    while i < total:
        line = lines[i]
        stripped = line.strip()

        # 代码块围栏
        if stripped.startswith('```'):
            if in_code_block:
                flush_body()
                _add_code_block(doc, code_lines, profile)
                code_lines = []
                in_code_block = False
            else:
                flush_body()
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            i += 1
            continue
        if in_code_block:
            code_lines.append(line.rstrip())
            i += 1
            continue

        if stripped.startswith('|'):
            flush_body()
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        if in_table:
            flush_table()

        if not stripped:
            flush_body()
            i += 1
            continue

        if re.fullmatch('[-*_]{3,}', stripped):
            flush_body()
            i += 1
            continue

        heading_match = re.match('^(#{1,6})\\s+(.+)$', stripped)
        if heading_match:
            flush_body()
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()

            if level in (1, 2):
                normalized = text.replace(' ', '').lower()
                if '参考文献' in normalized or 'references' in normalized:
                    in_references = True
                elif level <= 2:
                    in_references = False

            normalized_h = text.replace(' ', '').lower()
            is_abstract = (level == 2 and (
                normalized_h == '摘要'
                or normalized_h == 'abstract'
                or normalized_h == 'summary'
                or normalized_h == 'summarysheet'
                or normalized_h.startswith('摘要(')
                or normalized_h.startswith('摘要（')
            ))
            if is_abstract:
                in_abstract = True
                _add_abstract_heading(doc, text, profile)
                i += 1
                continue

            in_abstract = False
            is_first = (level == 1 and not first_h1_seen)
            if level == 1:
                first_h1_seen = True
            _add_heading(doc, text, level, profile, first_heading=is_first)
            i += 1
            continue

        img_match = re.match('^!\\[([^\\]]*)\\]\\(([^)]+)\\)\\s*$', stripped)
        if img_match:
            flush_body()
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            _add_image(doc, img_path, workspace, profile)
            # 图注：非空 alt 作为题注
            if alt_text and alt_text.strip():
                cap_para = doc.add_paragraph()
                cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_para.paragraph_format.first_line_indent = Pt(0)
                cap_run = cap_para.add_run(alt_text.strip())
                cap_run.font.size = Pt(10.5)
                _set_run_fonts(cap_run, profile, 'body')
            i += 1
            continue

        list_match = re.match('^(\\s*)([-*+]|\\d+[.)]) (.+)$', stripped)
        if list_match:
            flush_body()
            text = list_match.group(3)
            para = doc.add_paragraph()
            para.paragraph_format.first_line_indent = Pt(0)
            para.paragraph_format.left_indent = Cm(0.74)
            bullet = list_match.group(2)
            if re.match('\\d+', bullet):
                prefix = bullet + ' '
            else:
                prefix = '• '
            run = para.add_run(prefix + _clean_markdown_inline(text))
            run.font.size = Pt(profile.get('body', {}).get('font_size_pt', 12))
            _set_run_fonts(run, profile, 'body')
            i += 1
            continue

        body_buffer.append(stripped)
        i += 1

    flush_body()
    if in_code_block and code_lines:
        _add_code_block(doc, code_lines, profile)
    if in_table:
        flush_table()
    return


def main():
    parser = argparse.ArgumentParser(description='Markdown → DOCX 转换（中文学术格式）')
    parser.add_argument('--source', '-s', type=Path, required=True, help='源 Markdown 文件路径')
    parser.add_argument('--output', '-o', type=Path, required=True, help='输出 DOCX 文件路径')
    parser.add_argument('--profile', '-p', type=Path, default=None, help='样式配置 JSON 路径（默认使用 default_cn_thesis.json）')
    parser.add_argument('--workspace', '-w', type=Path, default=None, help='工作区根目录（用于解析图片相对路径）')
    parser.add_argument('--engine', '-e', type=str, default='auto', choices=['auto', 'python', 'node'], help='引擎选择：auto（自动）/ python（python-docx）/ node（docx.js）')
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

    try:
        result = markdown_to_docx(
            source_md=args.source,
            output_path=args.output,
            style_profile=args.profile,
            workspace=args.workspace,
            engine=args.engine,
        )
        print(f'✓ 导出成功: {result}')
        return
    except FileNotFoundError as e:
        print(f'✗ 错误: {e}', file=sys.stderr)
        sys.exit(1)
    except ValueError as e:
        print(f'✗ 编码错误: {e}', file=sys.stderr)
        sys.exit(2)
    except Exception as e:
        print(f'✗ 转换失败: {e}', file=sys.stderr)
        sys.exit(3)


if __name__ == '__main__':
    main()
