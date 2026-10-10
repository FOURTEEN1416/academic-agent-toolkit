"""docx 格式探针（skills/_utils/docx_format_probe.py）与 docx_export 页码页脚的定向测试。

覆盖轨2（docx 路线格式保障补齐）三态退出码 + 导出集成断言：
  - 干净 docx（PAGE 域页脚 + 空署名）→ 探针 rc=0；
  - 身份泄漏（core.creator 中文姓名 / 正文「指导教师：张三」）→ rc=1；
  - 无 docx → rc=2 且打印「不冒充通过」；
  - tools/docx_export.py 导出产物 zipfile 检查 word/footer*.xml 含 PAGE 域；
  - _utils 与 shared-scripts 逐字节孪生（md5 一致 + 同行为）。

夹具全部用 python-docx 运行时构造（禁止 heredoc 写含反斜杠夹具——本仓四次踩坑）。
注意 python-docx API：dc:creator 对应 core_properties.author（不存在 .creator 属性，
赋值会被静默丢弃——本测试编写时实测踩过）。
"""
import hashlib
import subprocess
import sys
import zipfile
from pathlib import Path

TOOLBOX = Path(__file__).resolve().parents[1]
PROBE = TOOLBOX / "skills" / "_utils" / "docx_format_probe.py"
PROBE_TWIN = TOOLBOX / "skills" / "shared-scripts" / "docx_format_probe.py"
DOCX_EXPORT = TOOLBOX / "tools" / "docx_export.py"


def _run_probe(directory: Path, script: Path = PROBE):
    return subprocess.run(
        [sys.executable, str(script), str(directory)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120,
    )


def _add_page_field_footer(doc):
    """给默认节页脚挂 PAGE 域（与 tools/docx_export.py 的 _add_page_number_footer 同构，
    测试独立构造，不反向 import 生产码，防生产码坏掉时测试跟着哑）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    para._p.append(fld)


def _make_docx(path: Path, *, page_footer: bool, author: str = "", body: str = "这是正文段落。"):
    from docx import Document

    doc = Document()
    doc.add_paragraph(body)
    if page_footer:
        _add_page_field_footer(doc)
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    doc.save(str(path))


# ---------- 三态退出码 ----------

def test_probe_clean_docx_exit_0(tmp_path):
    """干净夹具（PAGE 域页脚 + 空署名）→ rc=0，页码域 PASS。"""
    paper = tmp_path / "paper"
    paper.mkdir()
    _make_docx(paper / "main.docx", page_footer=True)
    proc = _run_probe(tmp_path)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "[页码域] PASS" in proc.stdout
    assert ": FAIL" not in proc.stdout  # 级别行形态「[检查名] FAIL:」；汇总行「无 FAIL」不算


def test_probe_identity_leak_in_core_exit_1(tmp_path):
    """core.creator=中文真实姓名样例 → rc=1。"""
    paper = tmp_path / "paper"
    paper.mkdir()
    _make_docx(paper / "main.docx", page_footer=True, author="张三")
    proc = _run_probe(tmp_path)
    assert proc.returncode == 1, proc.stdout
    assert "[身份泄漏] FAIL" in proc.stdout
    assert "张三" in proc.stdout


def test_probe_identity_leak_in_body_exit_1(tmp_path):
    """正文含「指导教师：张三」→ rc=1（匿名评审硬泄漏）。"""
    paper = tmp_path / "paper"
    paper.mkdir()
    _make_docx(paper / "main.docx", page_footer=True,
               body="本文方法见下节。指导教师：张三。")
    proc = _run_probe(tmp_path)
    assert proc.returncode == 1, proc.stdout
    assert "指导教师" in proc.stdout


def test_probe_no_docx_exit_2_no_fake_pass(tmp_path):
    """无 docx → rc=2 + 打印「不冒充通过」（SKIP 不冒充通过语义）。"""
    proc = _run_probe(tmp_path)
    assert proc.returncode == 2, proc.stdout
    assert "不冒充通过" in proc.stdout


# ---------- docx_export 集成：导出产物必带 PAGE 域页脚 ----------

def test_docx_export_footer_contains_page_field(tmp_path):
    """docx_export（python 引擎）导出的 docx：zipfile 检查 word/footer*.xml 含 PAGE 域，
    且产物过探针页码域检查（rc=0）。"""
    source = tmp_path / "paper.md"
    source.write_text("# 测试论文\n\n## 引言\n\n这是页码页脚集成测试的正文段落。\n",
                      encoding="utf-8", newline="\n")
    output = tmp_path / "paper" / "main.docx"
    output.parent.mkdir()

    proc = subprocess.run(
        [sys.executable, str(DOCX_EXPORT),
         "--source", str(source), "--output", str(output), "--engine", "python"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
    )
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-500:]
    assert output.is_file()

    with zipfile.ZipFile(output) as zf:
        footers = [n for n in zf.namelist() if n.startswith("word/footer") and n.endswith(".xml")]
        assert footers, "导出产物无 word/footer*.xml — 页码页脚未写入"
        joined = "".join(zf.read(n).decode("utf-8", errors="replace") for n in footers)
        assert 'w:instr="PAGE"' in joined or "<w:instrText" in joined and "PAGE" in joined, \
            "页脚 XML 无 PAGE 域指令"

    probe = _run_probe(output.parent)
    assert probe.returncode == 0, probe.stdout
    assert "[页码域] PASS" in probe.stdout


# ---------- 孪生纪律：_utils 与 shared-scripts 逐字节一致 ----------

def test_twin_copies_identical_md5_and_same_behavior(tmp_path):
    """两副本 md5 一致（硬规则：逐字节孪生），且 shared-scripts 副本跑同行为（无 docx → rc=2）。"""
    import filecmp

    assert PROBE.is_file() and PROBE_TWIN.is_file()
    assert filecmp.cmp(PROBE, PROBE_TWIN, shallow=False), "两副本内容不一致（逐字节孪生被破坏）"
    md5 = {hashlib.md5(p.read_bytes()).hexdigest() for p in (PROBE, PROBE_TWIN)}
    assert len(md5) == 1

    twin = _run_probe(tmp_path, script=PROBE_TWIN)
    utils = _run_probe(tmp_path, script=PROBE)
    assert twin.returncode == utils.returncode == 2
    assert twin.stdout == utils.stdout
